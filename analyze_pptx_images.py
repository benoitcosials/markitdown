import pptx
from pathlib import Path

# Load PPTX
pptx_files = list(Path('pptx-test').glob('Kickoff*.pptx'))
if not pptx_files:
    print("No PPTX file found")
    exit(1)

pptx_file = pptx_files[0]
print(f"Analyzing: {pptx_file.name}\n")

prs = pptx.Presentation(str(pptx_file))

print('=' * 80)
print('ANALYSE DÉTAILLÉE DES IMAGES PPTX')
print('=' * 80 + '\n')

for slide_num, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        # Check if picture
        is_picture = (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE or 
                     (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, 'image')))
        
        if is_picture and hasattr(shape, 'image'):
            img = shape.image
            size_kb = len(img.blob) / 1024
            
            print(f'Slide {slide_num} - {shape.name}')
            print(f'  Filename: {img.filename}')
            print(f'  Size: {len(img.blob):,} bytes ({size_kb:.1f} KB)')
            print(f'  Content-Type: {img.content_type}')
            print(f'  Dimensions (EMUs): {shape.width} x {shape.height}')
            print(f'  Dimensions (inches): {shape.width / 914400:.2f}" x {shape.height / 914400:.2f}"')
            
            # Check if placeholder
            if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER:
                print(f'  Type: PLACEHOLDER')
                try:
                    ph_type = shape.placeholder_format.type
                    print(f'  Placeholder Type: {ph_type}')
                except:
                    pass
            else:
                print(f'  Type: PICTURE')
            
            # Try to get more properties
            try:
                # Check XML element
                elem = shape._element
                
                # Check for description (alt text)
                try:
                    descr = shape._element._nvXxPr.cNvPr.attrib.get('descr', '')
                    if descr:
                        print(f'  Alt Text: {descr}')
                except:
                    pass
                
                # Check relationship ID (rId)
                try:
                    blip = shape._element.blipFill.blip
                    embed_id = blip.embed
                    print(f'  Embed ID: {embed_id}')
                except:
                    pass
                    
            except Exception as e:
                print(f'  XML Error: {e}')
            
            # Categorization heuristic
            if size_kb < 10:
                category = 'ICON (< 10 KB)'
            elif size_kb < 50:
                category = 'SMALL IMAGE (10-50 KB)'
            else:
                category = 'SCREENSHOT/PHOTO (> 50 KB)'
            
            print(f'  >>> CATEGORY: {category}')
            print()
