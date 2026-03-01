"""Debug slide 16 table and image structure."""
import sys

sys.path.insert(0, r"C:\Repos\markitdown\packages\markitdown\src")
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r".copilot-tracking\tests\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx"
prs = pptx.Presentation(pptx_path)
slide = prs.slides[15]  # 0-indexed, slide 16

print("=== Slide 16 Shapes Analysis ===\n")

for idx, shape in enumerate(slide.shapes):
    shape_type = shape.shape_type
    print(f"Shape {idx}: {shape_type.name if hasattr(shape_type, 'name') else shape_type}")
    print(f"  Position: left={shape.left}, top={shape.top}")
    print(f"  Size: width={shape.width}, height={shape.height}")
    
    if shape_type == MSO_SHAPE_TYPE.TABLE:
        table = shape.table
        print(f"  TABLE: {len(table.rows)}x{len(table.columns)}")
        print(f"  first_row (header): {table.first_row}")
        
        # Check a few cells for images
        print("  Checking cells for blipFill images...")
        for row_idx in range(min(3, len(table.rows))):
            for col_idx in range(len(table.columns)):
                cell = table.cell(row_idx, col_idx)
                cell_xml = cell._tc.xml if hasattr(cell, '_tc') else ""
                if "blip" in cell_xml.lower():
                    text = cell.text.strip()[:20] if cell.text else ""
                    print(f"    Cell[{row_idx},{col_idx}] has blip! text='{text}'")
                    
    elif shape_type == MSO_SHAPE_TYPE.PICTURE:
        print(f"  PICTURE: name='{shape.name}'")
        if hasattr(shape, 'image'):
            print(f"    content_type: {shape.image.content_type}")
            
    elif hasattr(shape, 'text') and shape.text:
        print(f"  Text: '{shape.text[:50]}...'")
    
    print()

# Check if images are placed OVER the table visually
print("\n=== Position Analysis ===")
table_shape = None
image_shapes = []

for shape in slide.shapes:
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        table_shape = shape
    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image_shapes.append(shape)

if table_shape:
    print("Table bounds:")
    print(f"  Left: {table_shape.left} to {table_shape.left + table_shape.width}")
    print(f"  Top: {table_shape.top} to {table_shape.top + table_shape.height}")
    
    print(f"\nImages ({len(image_shapes)} found):")
    for img in image_shapes:
        # Check if image is within table bounds
        in_table_x = table_shape.left <= img.left <= table_shape.left + table_shape.width
        in_table_y = table_shape.top <= img.top <= table_shape.top + table_shape.height
        
        print(f"  '{img.name}' at ({img.left}, {img.top}) - in_table: x={in_table_x}, y={in_table_y}")
