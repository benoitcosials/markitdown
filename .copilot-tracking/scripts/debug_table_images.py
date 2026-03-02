"""Debug script to analyze table and floating image positions in PPTX."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"C:\Repos\markitdown\.copilot-tracking\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx"
prs = Presentation(pptx_path)

for slide_idx, slide in enumerate(prs.slides, 1):
    if slide_idx not in (16, 17):
        continue
    
    print(f"\n{'='*80}")
    print(f"SLIDE {slide_idx}")
    print(f"{'='*80}")
    
    # Find tables and pictures
    tables = []
    pictures = []
    table_counter = 0
    
    for shape_idx, shape in enumerate(slide.shapes):
        if shape.has_table:
            tables.append((table_counter, shape))
            print(f"\nTABLE #{table_counter}: {shape.name} (shape_idx={shape_idx})")
            table_counter += 1
            print(f"  Position: left={shape.left}, top={shape.top}")
            print(f"  Size: width={shape.width}, height={shape.height}")
            
            tbl = shape.table
            print(f"  Dimensions: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
            
            # Print column boundaries
            col_left = shape.left
            print("\n  Column boundaries:")
            for col_idx, col in enumerate(tbl.columns):
                col_right = col_left + col.width
                print(f"    Col {col_idx}: left={col_left} -> right={col_right} (width={col.width})")
                col_left = col_right
            
            # Print row boundaries
            row_top = shape.top
            print("\n  Row boundaries:")
            for row_idx, row in enumerate(tbl.rows):
                row_bottom = row_top + row.height
                print(f"    Row {row_idx}: top={row_top} -> bottom={row_bottom} (height={row.height})")
                row_top = row_bottom
                
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(shape)
    
    # Print picture positions
    print(f"\nPICTURES ({len(pictures)} found):")
    for pic in pictures:
        cx = pic.left + pic.width // 2
        cy = pic.top + pic.height // 2
        print(f"  {pic.name}:")
        print(f"    Position: left={pic.left}, top={pic.top}")
        print(f"    Size: width={pic.width}, height={pic.height}")
        print(f"    Center: cx={cx}, cy={cy}")
        
        # Check which cell this image belongs to (for each table)
        for tbl_idx, tbl_shape in tables:
            tbl = tbl_shape.table
            tbl_left = tbl_shape.left
            tbl_top = tbl_shape.top
            tbl_right = tbl_left + tbl_shape.width
            tbl_bottom = tbl_top + tbl_shape.height
            
            # Check if center is within table bounds
            if tbl_left <= cx <= tbl_right and tbl_top <= cy <= tbl_bottom:
                # Find column
                col_left = tbl_shape.left
                found_col = -1
                for col_idx, col in enumerate(tbl.columns):
                    col_right = col_left + col.width
                    if col_left <= cx < col_right:
                        found_col = col_idx
                        break
                    col_left = col_right
                
                # Find row
                row_top = tbl_shape.top
                found_row = -1
                for row_idx, row in enumerate(tbl.rows):
                    row_bottom = row_top + row.height
                    if row_top <= cy < row_bottom:
                        found_row = row_idx
                        break
                    row_top = row_bottom
                
                print(f"    -> Mapped to cell ({found_row}, {found_col}) in TABLE #{tbl_idx} {tbl_shape.name}")
            else:
                print(f"    -> NOT in table {tbl_shape.name}")
