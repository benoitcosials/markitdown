"""Debug slide 10 to check bold formatting."""
from pathlib import Path

from pptx import Presentation

pptx_path = Path(r"C:\Repos\markitdown\.copilot-tracking\pptx\Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx")
prs = Presentation(pptx_path)

slide = prs.slides[9]  # Slide 10 (0-indexed)
print("=== Slide 10 ===\n")

for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    
    text_preview = shape.text_frame.text[:50].replace('\n', ' ')
    print(f"Shape {shape.shape_id}: {text_preview}...")
    
    for i, para in enumerate(shape.text_frame.paragraphs):
        para_text = para.text.strip()
        if not para_text:
            continue
            
        print(f"  Para {i} (level={para.level or 0}):")
        for j, run in enumerate(para.runs):
            if not run.text:
                continue
            bold = run.font.bold
            italic = run.font.italic
            print(f"    Run {j}: bold={bold}, italic={italic}, text='{run.text[:40]}'")
    print()
