"""
Script de diagnostic pour analyser le slide 19 du fichier PPTX.
"""
from pathlib import Path

import pptx

# Ouvrir le fichier PPTX
pptx_file = Path("pptx-test/z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx")
presentation = pptx.Presentation(pptx_file)

# Accéder au slide 19 (index 18, car 0-based)
slide_index = 18
slide = presentation.slides[slide_index]

print("=" * 80)
print(f"ANALYSE DU SLIDE {slide_index + 1}")
print("=" * 80)
print(f"\nNombre total de shapes: {len(slide.shapes)}")
print(f"\n{'#':<5} {'Type':<25} {'Nom':<40} {'Détails'}")
print("-" * 100)

for idx, shape in enumerate(slide.shapes):
    shape_type = str(type(shape).__name__)
    shape_name = shape.name
    
    # Détails spécifiques selon le type
    details = []
    
    # Vérifier si c'est une image
    if hasattr(shape, 'image'):
        try:
            details.append("IMAGE")
            details.append(f"Format: {shape.image.content_type}")
            details.append(f"Taille: {len(shape.image.blob)} bytes")
        except:
            details.append("IMAGE (erreur d'accès)")
    
    # Vérifier si c'est un groupe
    if hasattr(shape, 'shapes'):
        details.append(f"GROUPE ({len(shape.shapes)} shapes)")
    
    # Vérifier si c'est une forme
    if hasattr(shape, 'shape_type'):
        details.append(f"ShapeType: {shape.shape_type}")
    
    # Vérifier si ça a du texte
    if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
        if hasattr(shape, 'text') and shape.text.strip():
            text_preview = shape.text[:50].replace('\n', ' ')
            details.append(f"Texte: '{text_preview}...'")
    
    # Vérifier si c'est un graphique
    if hasattr(shape, 'has_chart') and shape.has_chart:
        details.append("CHART")
    
    # Vérifier si c'est une table
    if hasattr(shape, 'has_table') and shape.has_table:
        details.append("TABLE")
    
    # Position et taille
    if hasattr(shape, 'left') and hasattr(shape, 'top'):
        details.append(f"Pos: ({shape.left}, {shape.top})")
    
    if hasattr(shape, 'width') and hasattr(shape, 'height'):
        details.append(f"Size: {shape.width}x{shape.height}")
    
    details_str = " | ".join(details) if details else "Aucun détail"
    
    print(f"{idx:<5} {shape_type:<25} {shape_name:<40} {details_str}")

# Inspection détaillée des groupes
print(f"\n{'=' * 80}")
print("INSPECTION DES GROUPES")
print(f"{'=' * 80}\n")

for idx, shape in enumerate(slide.shapes):
    if hasattr(shape, 'shapes'):
        print(f"\n[Groupe #{idx}] {shape.name}")
        print(f"  Contient {len(shape.shapes)} sous-shapes:")
        for sub_idx, sub_shape in enumerate(shape.shapes):
            sub_type = str(type(sub_shape).__name__)
            sub_name = sub_shape.name
            sub_details = []
            
            if hasattr(sub_shape, 'image'):
                try:
                    sub_details.append(f"IMAGE - {sub_shape.image.content_type}")
                except:
                    sub_details.append("IMAGE (erreur)")
            
            if hasattr(sub_shape, 'shape_type'):
                sub_details.append(f"ShapeType: {sub_shape.shape_type}")
            
            sub_details_str = " | ".join(sub_details) if sub_details else ""
            print(f"    [{sub_idx}] {sub_type:<25} {sub_name:<35} {sub_details_str}")

print(f"\n{'=' * 80}")
