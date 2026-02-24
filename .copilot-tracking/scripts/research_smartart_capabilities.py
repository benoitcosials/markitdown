"""
Script d'exploration des capacités python-pptx pour SmartArt.

Objectif : Déterminer si python-pptx peut :
1. Détecter les SmartArt dans un PPTX
2. Extraire le texte et structure hiérarchique
3. Obtenir une image PNG du SmartArt

Date : 22 février 2026
Brief : BRIEF_05_SMARTART_EXTRACTION
"""

import sys
from pathlib import Path

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print("❌ python-pptx not installed. Run: pip install python-pptx")
    sys.exit(1)


def explore_smartart_capabilities(pptx_path: Path):
    """
    Explore SmartArt capabilities in a PPTX file.
    
    Args:
        pptx_path: Path to PPTX file to analyze
    """
    
    print(f"\n{'=' * 80}")
    print(f"Analyzing: {pptx_path.name}")
    print(f"{'=' * 80}")
    
    try:
        presentation = pptx.Presentation(pptx_path)
    except Exception as e:
        print(f"❌ Error loading PPTX: {e}")
        return
    
    smartart_found = False
    total_shapes = 0
    
    for slide_idx, slide in enumerate(presentation.slides, start=1):
        print(f"\n--- Slide {slide_idx} ({len(slide.shapes)} shapes) ---")
        
        for shape_idx, shape in enumerate(slide.shapes):
            total_shapes += 1
            
            # Get basic shape info
            shape_type = shape.shape_type
            shape_type_name = get_shape_type_name(shape_type)
            shape_name = shape.name
            
            # Check if shape might be SmartArt
            is_potential_smartart = detect_smartart(shape)
            
            if is_potential_smartart:
                smartart_found = True
                print(f"\n🎯 SMARTART DETECTED - Shape {shape_idx}")
                print(f"   Name: {shape_name}")
                print(f"   Type: {shape_type_name} ({shape_type})")
                print(f"   Class: {type(shape).__name__}")
                
                # Deep exploration
                explore_smartart_shape(shape, slide_idx, shape_idx)
            else:
                # Log non-SmartArt shapes briefly
                print(f"  Shape {shape_idx}: {shape_name[:40]:<40} | Type: {shape_type_name:<20} | {type(shape).__name__}")
    
    print(f"\n{'=' * 80}")
    print("Summary:")
    print(f"  Total shapes analyzed: {total_shapes}")
    print(f"  SmartArt found: {'✅ YES' if smartart_found else '❌ NO'}")
    print(f"{'=' * 80}\n")


def get_shape_type_name(shape_type: int) -> str:
    """Get human-readable name for shape type."""
    type_names = {
        1: "AUTO_SHAPE",
        2: "CALLOUT",
        3: "CHART",
        4: "COMMENT",
        5: "FREEFORM",
        6: "GROUP",
        7: "EMBEDDED_OLE",
        9: "LINE",
        10: "LINKED_OLE",
        11: "LINKED_PICTURE",
        12: "OLE_CONTROL",
        13: "PICTURE",
        14: "PLACEHOLDER",
        16: "MEDIA",
        17: "TEXT_BOX",
        18: "SCRIPT_ANCHOR",
        19: "TABLE",
        20: "CANVAS",
        21: "DIAGRAM",  # SMARTART ?
        22: "INK",
        23: "INK_COMMENT",
        24: "IGX_GRAPHIC",
        26: "WEB_VIDEO",
    }
    return type_names.get(shape_type, f"UNKNOWN_{shape_type}")


def detect_smartart(shape) -> bool:
    """
    Detect if a shape is a SmartArt.
    
    Returns:
        bool: True if shape is likely a SmartArt
    """
    
    # Method 1: Check shape type for DIAGRAM
    if hasattr(MSO_SHAPE_TYPE, 'DIAGRAM'):
        if shape.shape_type == MSO_SHAPE_TYPE.DIAGRAM:
            return True
    
    # Method 2: Check for smart_art attribute
    if hasattr(shape, 'smart_art'):
        if shape.smart_art is not None:
            return True
    
    # Method 3: Check for diagram attribute
    if hasattr(shape, 'diagram'):
        if shape.diagram is not None:
            return True
    
    # Method 4: Check XML for diagram URI
    try:
        xml_element = shape._element
        
        # SmartArt uses GraphicFrame
        if hasattr(xml_element, 'graphic'):
            graphic = xml_element.graphic
            if hasattr(graphic, 'graphicData'):
                graphic_data = graphic.graphicData
                uri = graphic_data.get('uri', '')
                
                # Diagram URI indicates SmartArt
                diagram_uri = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
                if diagram_uri in str(uri):
                    return True
    except:
        pass
    
    return False


def explore_smartart_shape(shape, slide_num: int, shape_num: int):
    """
    Deep dive exploration of SmartArt shape.
    
    Args:
        shape: python-pptx Shape object (suspected SmartArt)
        slide_num: Slide number
        shape_num: Shape number within slide
    """
    
    print(f"\n   {'─' * 70}")
    print(f"   SMARTART EXPLORATION - Slide {slide_num}, Shape {shape_num}")
    print(f"   {'─' * 70}")
    
    # 1. TEXT EXTRACTION
    print("\n   📝 TEXT EXTRACTION:")
    
    # Test 1a: shape.text
    if hasattr(shape, 'text'):
        try:
            text = shape.text
            print(f"      ✅ shape.text: '{text[:100]}...' ({len(text)} chars)")
        except:
            print("      ⚠️ shape.text exists but raises error")
    else:
        print("      ❌ No shape.text attribute")
    
    # Test 1b: shape.text_frame
    if hasattr(shape, 'text_frame'):
        try:
            text_frame = shape.text_frame
            if text_frame:
                text = text_frame.text
                print(f"      ✅ shape.text_frame.text: '{text[:100]}...'")
            else:
                print("      ⚠️ shape.text_frame is None")
        except:
            print("      ⚠️ shape.text_frame exists but error")
    else:
        print("      ❌ No text_frame attribute")
    
    # 2. SMARTART API
    print("\n   🔧 SMARTART API:")
    
    # Test 2a: shape.smart_art
    if hasattr(shape, 'smart_art'):
        smart_art = shape.smart_art
        print(f"      ✅ shape.smart_art exists: {smart_art}")
        
        if smart_art:
            print(f"         Type: {type(smart_art).__name__}")
            
            # List all public methods/attributes
            public_attrs = [attr for attr in dir(smart_art) if not attr.startswith('_')]
            print(f"         Public attributes: {public_attrs[:10]}...")  # First 10
            
            # Test common API patterns
            for attr_name in ['nodes', 'points', 'all_nodes', 'node_list', 'items']:
                if hasattr(smart_art, attr_name):
                    try:
                        attr_value = getattr(smart_art, attr_name)
                        print(f"         ✅ {attr_name}: {attr_value}")
                    except:
                        print(f"         ⚠️ {attr_name} exists but error accessing")
    else:
        print("      ❌ No smart_art attribute")
    
    # Test 2b: shape.diagram
    if hasattr(shape, 'diagram'):
        diagram = shape.diagram
        print(f"      ✅ shape.diagram exists: {diagram}")
        if diagram:
            print(f"         Type: {type(diagram).__name__}")
    else:
        print("      ❌ No diagram attribute")
    
    # 3. IMAGE EXTRACTION
    print("\n   🖼️ IMAGE EXTRACTION:")
    
    # Test 3a: shape.image
    if hasattr(shape, 'image'):
        try:
            image = shape.image
            if image:
                blob_size = len(image.blob)
                content_type = image.content_type
                print(f"      ✅ shape.image.blob: {blob_size} bytes")
                print(f"         Content-Type: {content_type}")
            else:
                print("      ⚠️ shape.image is None")
        except:
            print("      ⚠️ shape.image exists but error")
    else:
        print("      ❌ No image attribute (SmartArt not stored as image)")
    
    # 4. XML STRUCTURE
    print("\n   📄 XML STRUCTURE:")
    
    try:
        xml_element = shape._element
        print(f"      ✅ XML element tag: {xml_element.tag}")
        
        # Explore graphic element
        if hasattr(xml_element, 'graphic'):
            graphic = xml_element.graphic
            print("      ✅ Has <graphic> element")
            
            if hasattr(graphic, 'graphicData'):
                graphic_data = graphic.graphicData
                uri = graphic_data.get('uri', '')
                print(f"         URI: {uri}")
                
                # Try to find diagram parts
                if 'diagram' in uri:
                    print("         🎯 CONFIRMED: Diagram URI detected!")
                    explore_diagram_xml(shape)
        else:
            print("      ❌ No <graphic> element")
            
    except Exception as e:
        print(f"      ⚠️ Error exploring XML: {e}")
    
    print(f"   {'─' * 70}\n")


def explore_diagram_xml(shape):
    """
    Explore diagram XML structure to find text nodes.
    
    Args:
        shape: SmartArt shape
    """
    
    print("\n      📊 DIAGRAM XML EXPLORATION:")
    
    try:
        # Access part relationships to find diagram data
        # SmartArt consists of multiple parts: data, layout, colors, style
        
        # Get shape element
        graphic_frame = shape._element
        
        # Navigate to diagram reference
        if hasattr(graphic_frame, 'graphic'):
            graphic = graphic_frame.graphic
            
            # Check for relIds (relationships to diagram parts)
            nsmap = graphic.nsmap
            
            # Try to find diagram data part
            # This is complex - diagram data is in a separate part file
            print(f"         XML Namespaces: {list(nsmap.keys())}")
            
            # List all attributes
            for key, value in graphic.attrib.items():
                print(f"         Attribute {key}: {value}")
            
            # Try to access diagram parts through shape.part
            if hasattr(shape, 'part'):
                part = shape.part
                print(f"         ✅ Shape has .part: {part}")
                
                if hasattr(part, 'related_parts'):
                    related = part.related_parts
                    print(f"         ✅ Related parts: {len(related)} parts")
                    
                    # Look for diagram parts
                    for rel_id, related_part in related.items():
                        part_name = related_part.partname if hasattr(related_part, 'partname') else 'unknown'
                        print(f"            - {rel_id}: {part_name}")
                        
                        # Check if it's a diagram data part
                        if 'diagram' in str(part_name).lower():
                            print("              🎯 Found diagram part!")
                            try:
                                # Try to read diagram XML
                                diagram_xml = related_part.blob
                                print(f"              Diagram XML size: {len(diagram_xml)} bytes")
                                
                                # Try parsing with lxml
                                try:
                                    from lxml import etree
                                    root = etree.fromstring(diagram_xml)
                                    print(f"              ✅ Parsed XML root: {root.tag}")
                                    
                                    # Look for text nodes
                                    # Namespace for diagram: dgm
                                    dgm_ns = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
                                    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
                                    
                                    # Find all <dgm:pt> (points/nodes)
                                    points = root.findall(f'.//{{{dgm_ns}}}pt')
                                    print(f"              ✅ Found {len(points)} <dgm:pt> nodes")
                                    
                                    # Extract text from first few nodes
                                    for idx, pt in enumerate(points[:5]):  # First 5 nodes
                                        # Find text <a:t>
                                        text_elem = pt.find(f'.//{{{a_ns}}}t')
                                        if text_elem is not None and text_elem.text:
                                            print(f"                 Node {idx}: '{text_elem.text}'")
                                    
                                except ImportError:
                                    print("              ⚠️ lxml not available, can't parse XML")
                                except Exception as xml_error:
                                    print(f"              ⚠️ XML parsing error: {xml_error}")
                                
                            except Exception as e:
                                print(f"              ⚠️ Error reading diagram part: {e}")
            
    except Exception as e:
        print(f"         ⚠️ Error exploring diagram XML: {e}")


def main():
    """Run SmartArt exploration on test files."""
    
    print("\n" + "=" * 80)
    print("SMARTART CAPABILITIES RESEARCH - python-pptx")
    print("=" * 80)
    
    # Test files to analyze
    test_files = [
        # Existing test files in pptx-test/
        "pptx-test/z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx",
        "pptx-test/Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx",
        "pptx-test/Standard QA - Design des essais.pptx",
        
        # Add test file with known SmartArt if available
        # "test_smartart.pptx",
    ]
    
    files_analyzed = 0
    
    for pptx_file in test_files:
        path = Path(pptx_file)
        if path.exists():
            explore_smartart_capabilities(path)
            files_analyzed += 1
        else:
            print(f"⚠️ File not found: {pptx_file}")
    
    if files_analyzed == 0:
        print("\n❌ No PPTX files found to analyze.")
        print("\n💡 To test SmartArt detection:")
        print("   1. Create a PPTX with SmartArt in PowerPoint/LibreOffice")
        print("   2. Save it in pptx-test/ folder")
        print("   3. Add the filename to this script")
        print("   4. Run this script again")
    else:
        print(f"\n✅ Analysis complete. {files_analyzed} file(s) analyzed.")
        print("\n📝 Next steps:")
        print("   1. Review the output above")
        print("   2. Document findings in research file")
        print("   3. Update BRIEF_05 based on capabilities discovered")


if __name__ == "__main__":
    main()
