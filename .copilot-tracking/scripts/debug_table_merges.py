"""Debug script to analyze table cell merges in PPTX."""
from pptx import Presentation
from pptx.oxml.ns import qn

pptx_path = r"C:\Repos\markitdown\.copilot-tracking\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx"
prs = Presentation(pptx_path)

for slide_idx, slide in enumerate(prs.slides, 1):
    if slide_idx != 16:
        continue
    
    print(f"\n{'='*80}")
    print(f"SLIDE {slide_idx}")
    print(f"{'='*80}")
    
    for shape in slide.shapes:
        if not shape.has_table:
            continue
            
        tbl = shape.table
        print(f"\nTABLE: {shape.name}")
        print(f"Dimensions: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
        
        # Print each cell with merge info
        for row_idx, row in enumerate(tbl.rows):
            print(f"\n  Row {row_idx}:")
            for col_idx, cell in enumerate(row.cells):
                # Get cell text (truncated)
                text = cell.text[:30].replace('\n', ' ') if cell.text else "(empty)"
                
                # Check for merge attributes in XML
                tc = cell._tc  # Access underlying XML element
                gridSpan = tc.get(qn('a:gridSpan'))
                rowSpan = tc.get(qn('a:rowSpan'))
                hMerge = tc.get(qn('a:hMerge'))
                vMerge = tc.get(qn('a:vMerge'))
                
                merge_info = ""
                if gridSpan and int(gridSpan) > 1:
                    merge_info += f" gridSpan={gridSpan}"
                if rowSpan and int(rowSpan) > 1:
                    merge_info += f" rowSpan={rowSpan}"
                if hMerge:
                    merge_info += " hMerge"
                if vMerge:
                    merge_info += " vMerge"
                
                # Check is_spanned attribute
                spanned = ""
                if hasattr(cell, 'is_spanned') and cell.is_spanned:
                    spanned = " [SPANNED]"
                
                print(f"    Col {col_idx}: '{text}'{merge_info}{spanned}")
