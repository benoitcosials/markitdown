"""Debug script to analyze table cell colors in PPTX."""
import sys

sys.path.insert(0, r"C:\Repos\markitdown\packages\markitdown\src")

import pptx
from pptx.enum.dml import MSO_FILL_TYPE

# Load presentation
pptx_path = r".copilot-tracking\tests\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx"
prs = pptx.Presentation(pptx_path)

# Analyze slide 2
slide = prs.slides[1]  # 0-indexed
print("=== Slide 2 Analysis ===\n")

for shape in slide.shapes:
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.TABLE:
        table = shape.table
        print(f"Table found: {len(table.rows)} rows x {len(table.columns)} cols\n")
        
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                text = cell.text.strip()[:20] if cell.text else ""
                
                # Check fill
                try:
                    fill = cell.fill
                    fill_type = fill.type
                    
                    if fill_type is not None:
                        print(f"Cell [{row_idx},{col_idx}]: fill_type={fill_type}")
                        
                        if fill_type == MSO_FILL_TYPE.SOLID:
                            fore_color = fill.fore_color
                            print(f"  -> fore_color.type = {fore_color.type}")
                            if hasattr(fore_color, 'rgb') and fore_color.rgb:
                                print(f"  -> RGB = {fore_color.rgb}")
                            if hasattr(fore_color, 'theme_color'):
                                print(f"  -> theme_color = {fore_color.theme_color}")
                        
                        print(f"  -> text = '{text}'")
                except Exception as e:
                    print(f"Cell [{row_idx},{col_idx}]: Error - {e}")

print("\n=== Text Encoding Check ===\n")
# Check raw text encoding
for shape in slide.shapes:
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.TABLE:
        cell = shape.table.cell(0, 0)
        text = cell.text
        print(f"Raw text: {repr(text)}")
        print(f"Bytes: {text.encode('utf-8')}")
        break
