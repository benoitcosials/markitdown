"""Debug slide 7 to analyze XML structure and algorithm decisions."""
from pathlib import Path

from lxml import etree
from pptx import Presentation

pptx_path = Path(r"C:\Repos\markitdown\.copilot-tracking\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx")
prs = Presentation(pptx_path)

slide = prs.slides[6]  # Slide 7 (0-indexed)
slide_height = prs.slide_height
slide_width = prs.slide_width

print("=== Slide 7 Analysis ===")
print(f"Slide dimensions: {slide_width} x {slide_height} EMU\n")

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

for shape in slide.shapes:
    if not hasattr(shape, 'text_frame') or not shape.has_text_frame:
        continue
    
    y = shape.top or 0
    x = shape.left or 0
    
    print("=" * 80)
    print(f"Shape: {shape.name} (ID: {shape.shape_id})")
    print(f"Position: Y={y} ({y/914400:.2f}in), X={x} ({x/914400:.2f}in)")
    print(f"Text preview: '{shape.text_frame.text[:60].replace(chr(10), ' ')}...'")
    
    # Check if hierarchical
    is_hierarchical = any((p.level or 0) > 0 for p in shape.text_frame.paragraphs)
    print(f"is_hierarchical: {is_hierarchical}")
    print()
    
    for i, para in enumerate(shape.text_frame.paragraphs):
        para_text = para.text.strip()
        if not para_text:
            continue
        
        level = para.level or 0
        
        # Get font size
        font_size = None
        if para.runs:
            run = para.runs[0]
            if run.font.size:
                font_size = run.font.size.pt
        
        # Check buNone in XML
        has_buNone = False
        has_buChar = False
        has_buAutoNum = False
        bullet_info = "none"
        
        pPr = para._p.find('.//a:pPr', NS)
        if pPr is not None:
            buNone = pPr.find('.//a:buNone', NS)
            buChar = pPr.find('.//a:buChar', NS)
            buAutoNum = pPr.find('.//a:buAutoNum', NS)
            has_buNone = buNone is not None
            has_buChar = buChar is not None
            has_buAutoNum = buAutoNum is not None
            
            if has_buNone:
                bullet_info = "buNone"
            elif has_buChar:
                char = buChar.get('char', '?')
                bullet_info = f"buChar('{char}')"
            elif has_buAutoNum:
                bullet_info = "buAutoNum"
        
        # Algorithm decision
        has_no_bullet = has_buNone
        is_bullet = level > 0 or (is_hierarchical and not has_no_bullet)
        
        print(f"  Para {i}: level={level}, size={font_size}pt")
        print(f"    XML bullet: {bullet_info}")
        print(f"    Algorithm: is_bullet={is_bullet} (is_hierarchical={is_hierarchical}, has_no_bullet={has_no_bullet})")
        print(f"    Text: '{para_text[:70]}'")
        print()
    
    # Show raw XML for first few paragraphs
    print("  --- Raw XML (first 2 paragraphs) ---")
    for i, para in enumerate(shape.text_frame.paragraphs[:2]):
        if para.text.strip():
            xml_str = etree.tostring(para._p, pretty_print=True, encoding='unicode')
            # Truncate long XML
            if len(xml_str) > 800:
                xml_str = xml_str[:800] + "\n    ... (truncated)"
            print(f"  Para {i}:")
            print(xml_str)
