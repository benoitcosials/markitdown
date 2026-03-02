"""Debug slide 4 to analyze shape positions."""
from pathlib import Path

from pptx import Presentation

pptx_path = Path(r"C:\Repos\markitdown\.copilot-tracking\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx")
prs = Presentation(pptx_path)

slide = prs.slides[3]  # Slide 4 (0-indexed)
slide_height = prs.slide_height

print("=== Slide 4 Position Analysis ===")
print(f"Slide height: {slide_height} EMU ({slide_height / 914400:.1f} inches)\n")

shapes_info = []

for shape in slide.shapes:
    y = shape.top or 0
    x = shape.left or 0
    
    # Get shape type/content
    if hasattr(shape, 'text_frame') and shape.has_text_frame:
        text = shape.text_frame.text[:50].replace('\n', ' ')
        shape_type = 'text'
    elif shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        shape_type = 'group/smartart'
        text = "(SmartArt/Group)"
    elif hasattr(shape, 'image'):
        shape_type = 'image'
        text = "(image)"
    else:
        shape_type = str(shape.shape_type)
        text = ""
    
    shapes_info.append({
        'id': shape.shape_id,
        'name': shape.name,
        'y': y,
        'x': x,
        'y_inches': y / 914400,
        'x_inches': x / 914400,
        'type': shape_type,
        'text': text
    })

# Sort by Y then X (simple sort)
shapes_info.sort(key=lambda s: (s['y'], s['x']))

print("Shapes sorted by Y, then X:")
print("-" * 80)
for s in shapes_info:
    y_band = int((s['y'] / slide_height) * 20)
    print(f"Y={s['y']:>8} ({s['y_inches']:>5.2f}in) band={y_band:>2} | X={s['x']:>8} | {s['type']:>12} | {s['name'][:30]}")
    if s['text']:
        print(f"    -> '{s['text']}'")

print("\n" + "=" * 80)
print("Current algorithm uses Y-bands (5% tolerance):")
print("-" * 80)

# Show what the current algorithm does
shapes_info_banded = sorted(shapes_info, key=lambda s: (int((s['y'] / slide_height) * 20), s['x']))
for s in shapes_info_banded:
    y_band = int((s['y'] / slide_height) * 20)
    print(f"band={y_band:>2} Y={s['y']:>8} X={s['x']:>8} | {s['name'][:30]}")
