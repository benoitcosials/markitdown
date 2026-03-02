"""Debug slide 11 to analyze heading detection algorithm."""
from pathlib import Path

from pptx import Presentation

pptx_path = Path(r"C:\Repos\markitdown\.copilot-tracking\pptx\Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx")
prs = Presentation(pptx_path)

slide = prs.slides[10]  # Slide 11 (0-indexed)
print("=== Slide 11 Analysis ===\n")

# Collect all paragraph sizes
all_sizes = []

for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    
    shape_name = shape.name if hasattr(shape, 'name') else f"Shape {shape.shape_id}"
    text_preview = shape.text_frame.text[:40].replace('\n', ' ')
    print(f"--- {shape_name} (ID:{shape.shape_id}): '{text_preview}...'")
    
    # Check if hierarchical
    is_hierarchical = any((p.level or 0) > 0 for p in shape.text_frame.paragraphs)
    print(f"    is_hierarchical: {is_hierarchical}")
    
    for i, para in enumerate(shape.text_frame.paragraphs):
        para_text = para.text.strip()
        if not para_text:
            continue
        
        # Get font size from first run
        font_size = None
        if para.runs:
            run = para.runs[0]
            if run.font.size:
                font_size = run.font.size.pt
        
        level = para.level or 0
        
        # Check buNone
        has_buNone = False
        try:
            pPr = para._p.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
            if pPr is not None:
                buNone = pPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}buNone')
                has_buNone = buNone is not None
        except:
            pass
        
        all_sizes.append((font_size, len(para_text), para_text[:50]))
        
        print(f"    Para {i}: level={level}, size={font_size}pt, buNone={has_buNone}")
        print(f"             text='{para_text[:60]}'")
    print()

# Analyze heading detection logic
print("=== Heading Detection Analysis ===")
print("\nAll paragraph sizes collected:")
for size, char_count, text in all_sizes:
    print(f"  {size}pt ({char_count} chars): '{text}'")

# Find dominant size (most characters)
size_char_count = {}
for size, char_count, _ in all_sizes:
    if size:
        size_char_count[size] = size_char_count.get(size, 0) + char_count

print("\nCharacter count per size:")
for size, count in sorted(size_char_count.items(), key=lambda x: -x[1]):
    print(f"  {size}pt: {count} chars")

if size_char_count:
    dominant_size = max(size_char_count.items(), key=lambda x: x[1])[0]
    print(f"\nDominant size: {dominant_size}pt")
    
    heading_sizes = sorted([s for s in size_char_count.keys() if s > dominant_size], reverse=True)
    print(f"Heading sizes (> dominant): {heading_sizes}")
