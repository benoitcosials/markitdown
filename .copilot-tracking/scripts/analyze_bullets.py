"""Analyze bullet properties of slide 6 paragraphs."""
from pptx import Presentation

pptx_path = r'.copilot-tracking\pptx\Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx'
prs = Presentation(pptx_path)

slide = prs.slides[5]
print('=== Slide 6 shapes ===')
for shape in slide.shapes:
    if shape.has_text_frame:
        text = shape.text[:50].replace(chr(11), ' ').replace('\n', ' ')
        print(f'Shape: {text}...')
        for i, para in enumerate(shape.text_frame.paragraphs):
            pt = para.text[:40].replace(chr(11), ' ')
            level = para.level or 0
            bullet_info = 'unknown'
            try:
                pPr = para._p.pPr
                if pPr is not None:
                    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
                    buNone = pPr.find(f'.//{ns}buNone')
                    buChar = pPr.find(f'.//{ns}buChar')
                    buAutoNum = pPr.find(f'.//{ns}buAutoNum')
                    if buNone is not None:
                        bullet_info = 'buNone (NO BULLET)'
                    elif buChar is not None:
                        char = buChar.get('char')
                        bullet_info = f'buChar={char}'
                    elif buAutoNum is not None:
                        bullet_info = 'buAutoNum'
                    else:
                        bullet_info = 'no-explicit-bullet-prop'
                else:
                    bullet_info = 'pPr=None'
            except Exception as e:
                bullet_info = f'error: {e}'
            print(f'  P{i} level={level} {bullet_info}: {pt}')
        print()
