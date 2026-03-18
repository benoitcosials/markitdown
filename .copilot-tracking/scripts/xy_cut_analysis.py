import pptx

DGM = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'

def is_sa(s):
    try:
        if not hasattr(s, '_element') or not hasattr(s._element, 'graphic'):
            return False
        return DGM in s._element.graphic.graphicData.get('uri', '')
    except:
        return False

p = pptx.Presentation('.copilot-tracking/pptx/Standard QA - Concept de base.pptx')
sh = p.slide_height
sw = p.slide_width

for idx, slide in enumerate(p.slides):
    ts = slide.shapes.title
    shapes_data = []
    for s in slide.shapes:
        if s == ts:
            continue
        t = s.top or 0
        l = s.left or 0
        h = s.height or 0
        w = s.width or 0
        if is_sa(s):
            tp = 'SA'
        elif s.has_table:
            tp = 'TBL'
        elif s.has_text_frame:
            tp = 'TXT'
        elif hasattr(s, 'image'):
            tp = 'PIC'
        else:
            continue
        tx = ''
        try:
            tx = s.text_frame.text[:25].replace('\n', ' | ')
        except:
            pass
        shapes_data.append((l, l + w, t, t + h, tp, tx))
    if len(shapes_data) < 2:
        continue

    print(f'=== Slide {idx+1} ({len(shapes_data)} shapes) ===')
    for xl, xr, yt, yb, tp, tx in shapes_data:
        print(f'  {tp:<4} L={xl/sw*100:5.1f}% R={xr/sw*100:5.1f}% T={yt/sh*100:5.1f}% B={yb/sh*100:5.1f}%  [{tx}]')

    # V-gaps
    events = []
    for xl, xr, yt, yb, tp, tx in shapes_data:
        events.append((xl, 1))
        events.append((xr, -1))
    events.sort()
    depth = 0
    prev = 0
    vgaps = []
    for x_val, d in events:
        if depth == 0 and x_val > prev and prev > 0:
            gw = (x_val - prev) / sw * 100
            if gw > 1.0:
                vgaps.append((prev / sw * 100, x_val / sw * 100, gw))
        depth += d
        prev = x_val

    # H-gaps
    events = []
    for xl, xr, yt, yb, tp, tx in shapes_data:
        events.append((yt, 1))
        events.append((yb, -1))
    events.sort()
    depth = 0
    prev = 0
    hgaps = []
    for y_val, d in events:
        if depth == 0 and y_val > prev and prev > 0:
            gh = (y_val - prev) / sh * 100
            if gh > 1.0:
                hgaps.append((prev / sh * 100, y_val / sh * 100, gh))
        depth += d
        prev = y_val

    if vgaps:
        print(f'  V-gaps: {[f"{g[0]:.1f}-{g[1]:.1f}% ({g[2]:.1f}%)" for g in vgaps]}')
    if hgaps:
        print(f'  H-gaps: {[f"{g[0]:.1f}-{g[1]:.1f}% ({g[2]:.1f}%)" for g in hgaps]}')
    if not vgaps and not hgaps:
        print('  No gaps > 1%')
    print()
