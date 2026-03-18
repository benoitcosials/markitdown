"""Analyze Slide 4 shape positions to debug X-Y Cut ordering."""

import pptx

DGM = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'

def is_sa(s):
    try:
        if not hasattr(s, '_element') or not hasattr(s._element, 'graphic'):
            return False
        return DGM in s._element.graphic.graphicData.get('uri', '')
    except Exception:
        return False

p = pptx.Presentation(
    '.copilot-tracking/pptx/z 015_ARCH-00xx_WS4_Bancaire'
    '-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx'
)
sh = p.slide_height
sw = p.slide_width

slide = p.slides[3]  # Slide 4 (0-indexed)
ts = slide.shapes.title

print(f"Slide dimensions: {sw} x {sh}")
print(f"Title: {ts.text[:40] if ts else 'None'}")
print()

for s in slide.shapes:
    if s == ts:
        print(f"  [TITLE] top={s.top} left={s.left} "
              f"h={s.height} w={s.width}")
        continue
    t = s.top or 0
    l = s.left or 0
    h = s.height or 0
    w = s.width or 0

    if is_sa(s):
        tp = 'SA'
    elif s.has_table:
        tp = 'TBL'
    elif s.has_text_frame:
        tp = 'TXT'
        txt = s.text_frame.text.strip()
        if not txt:
            tp = 'TXT_EMPTY'
    elif hasattr(s, 'image'):
        tp = 'PIC'
    else:
        tp = f'OTHER({s.shape_type})'

    tx = ''
    try:
        tx = s.text_frame.text[:40].replace('\n', ' | ')
    except Exception:
        tx = '...'

    print(f"  [{tp}] top={t} ({t/sh*100:.1f}%) "
          f"left={l} ({l/sw*100:.1f}%) "
          f"h={h} ({h/sh*100:.1f}%) "
          f"w={w} ({w/sw*100:.1f}%) "
          f"bot={(t+h)/sh*100:.1f}% "
          f"right={(l+w)/sw*100:.1f}%")
    print(f"         text=[{tx}]")

print()

# Now simulate the X-Y Cut algorithm
print("=== Simulating X-Y Cut ===")
shapes_data = []
for s in slide.shapes:
    if s == ts:
        continue
    t = s.top or 0
    l = s.left or 0
    h = s.height or 0
    w = s.width or 0

    if is_sa(s):
        tp = 'SA'
    elif s.has_table:
        tp = 'TBL'
    elif s.has_text_frame:
        txt = s.text_frame.text.strip()
        if not txt:
            continue  # filtered out
        tp = 'TXT'
    elif hasattr(s, 'image'):
        tp = 'PIC'
    else:
        continue

    tx = ''
    try:
        tx = s.text_frame.text[:30].replace('\n', ' | ')
    except Exception:
        tx = '...'

    shapes_data.append({
        'type': tp, 'text': tx,
        'left': l, 'right': l + w,
        'top': t, 'bottom': t + h,
    })

print(f"\nFiltered shapes: {len(shapes_data)}")
for sd in shapes_data:
    print(f"  {sd['type']:<4} "
          f"L={sd['left']/sw*100:.1f}% R={sd['right']/sw*100:.1f}% "
          f"T={sd['top']/sh*100:.1f}% B={sd['bottom']/sh*100:.1f}% "
          f"[{sd['text']}]")

# Find V-gaps
events = []
for sd in shapes_data:
    events.append((sd['left'], 1))
    events.append((sd['right'], -1))
events.sort()
depth = 0
prev = 0
vgaps = []
for x_val, d in events:
    if depth == 0 and x_val > prev and prev > 0:
        gw = (x_val - prev) / sw * 100
        if gw > 0.5:
            vgaps.append((prev / sw * 100, x_val / sw * 100, gw))
    depth += d
    prev = x_val
print(f"\nV-gaps: {[f'{g[0]:.1f}-{g[1]:.1f}% ({g[2]:.1f}%)' for g in vgaps]}")

# Find H-gaps
events = []
for sd in shapes_data:
    events.append((sd['top'], 1))
    events.append((sd['bottom'], -1))
events.sort()
depth = 0
prev = 0
hgaps = []
for y_val, d in events:
    if depth == 0 and y_val > prev and prev > 0:
        gh = (y_val - prev) / sh * 100
        if gh > 0.5:
            hgaps.append((prev / sh * 100, y_val / sh * 100, gh))
    depth += d
    prev = y_val
print(f"H-gaps: {[f'{g[0]:.1f}-{g[1]:.1f}% ({g[2]:.1f}%)' for g in hgaps]}")

# Show what the X-Y Cut algorithm would do
if vgaps:
    best_v = max(vgaps, key=lambda g: g[2])
    print(f"\nBest V-gap: {best_v[0]:.1f}-{best_v[1]:.1f}% "
          f"(ratio={best_v[2]:.1f}%)")
if hgaps:
    best_h = max(hgaps, key=lambda g: g[2])
    print(f"Best H-gap: {best_h[0]:.1f}-{best_h[1]:.1f}% "
          f"(ratio={best_h[2]:.1f}%)")
