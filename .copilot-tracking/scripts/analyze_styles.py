"""Analyze text styles and hyperlinks in PPTX."""
from pptx import Presentation

pptx_path = r'.copilot-tracking\pptx\Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx'
prs = Presentation(pptx_path)

# Find shapes with interesting formatting
for slide_idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    font = run.font
                    has_style = (font.bold or font.italic or font.underline or 
                                 getattr(font, 'strikethrough', None))
                    has_link = run.hyperlink and run.hyperlink.address
                    
                    if has_style or has_link:
                        text = run.text[:30] if run.text else ''
                        print(f'Slide {slide_idx+1}: "{text}"')
                        print(f'  bold={font.bold}, italic={font.italic}')
                        print(f'  underline={font.underline}')
                        print(f'  strikethrough={getattr(font, "strikethrough", None)}')
                        if has_link:
                            print(f'  hyperlink={run.hyperlink.address}')
                        print()
