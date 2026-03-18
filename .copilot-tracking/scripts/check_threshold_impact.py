"""Check impact of lowering gap threshold from 2% to 1.5%."""

import pptx

DGM = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'

def is_sa(s):
    try:
        if not hasattr(s, '_element') or not hasattr(s._element, 'graphic'):
            return False
        return DGM in s._element.graphic.graphicData.get('uri', '')
    except Exception:
        return False

files = [
    ('.copilot-tracking/pptx/Standard QA - Concept de base.pptx', 'QA'),
    ('.copilot-tracking/pptx/z 015_ARCH-00xx_WS4_Bancaire'
     '-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx', 'Bancaire'),
]

for fpath, label in files:
    p = pptx.Presentation(fpath)
    sh = p.slide_height
    sw = p.slide_width
    print(f"\n{'='*60}")
    print(f"  {label}")
    print(f"{'='*60}")

    for idx, slide in enumerate(p.slides):
        ts = slide.shapes.title
        shapes_data = []
        for s in slide.shapes:
            if s == ts:
                continue
            t = s.top or 0
            l = s.left or 0
            h = s.height or 0
            w = s.width or 0
            if is_sa(s):
                tp = 'SA'
            elif s.has_table:
                tp = 'TBL'
            elif s.has_text_frame:
                if not s.text_frame.text.strip():
                    continue
                tp = 'TXT'
            elif hasattr(s, 'image'):
                tp = 'PIC'
            else:
                continue
            shapes_data.append((l, l + w, t, t + h, tp))
        if len(shapes_data) < 2:
            continue

        # V-gaps
        events = []
        for xl, xr, yt, yb, tp in shapes_data:
            events.append((xl, 1))
            events.append((xr, -1))
        events.sort()
        depth = 0
        prev = 0
        vgaps = []
        for x_val, d in events:
            if depth == 0 and x_val > prev and prev > 0:
                gw = (x_val - prev) / sw * 100
                if gw > 1.0:
                    vgaps.append(gw)
            depth += d
            prev = x_val

        # H-gaps
        events = []
        for xl, xr, yt, yb, tp in shapes_data:
            events.append((yt, 1))
            events.append((yb, -1))
        events.sort()
        depth = 0
        prev = 0
        hgaps = []
        for y_val, d in events:
            if depth == 0 and y_val > prev and prev > 0:
                gh = (y_val - prev) / sh * 100
                if gh > 1.0:
                    hgaps.append(gh)
            depth += d
            prev = y_val

        # Report gaps in the 1.5%-2.0% range (new vs old threshold)
        new_vgaps = [g for g in vgaps if 1.5 <= g < 2.0]
        new_hgaps = [g for g in hgaps if 1.5 <= g < 2.0]
        if new_vgaps or new_hgaps:
            types = [tp for _, _, _, _, tp in shapes_data]
            print(f"  Slide {idx+1} ({len(shapes_data)} shapes, "
                  f"types={types}):")
            if new_vgaps:
                print(f"    NEW V-gaps (1.5-2%): "
                      f"{[f'{g:.1f}%' for g in new_vgaps]}")
            if new_hgaps:
                print(f"    NEW H-gaps (1.5-2%): "
                      f"{[f'{g:.1f}%' for g in new_hgaps]}")

    print("  (Done)")
