"""Debug slide 16: calculate cell positions and map images to cells."""
import sys

sys.path.insert(0, r"C:\Repos\markitdown\packages\markitdown\src")
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r".copilot-tracking\tests\pptx\z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx"
prs = pptx.Presentation(pptx_path)
slide = prs.slides[15]  # Slide 16

# Find table and images
table_shape = None
image_shapes = []

for shape in slide.shapes:
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        table_shape = shape
    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image_shapes.append(shape)

if not table_shape:
    print("No table found")
    sys.exit(1)

table = table_shape.table
print(f"Table: {len(table.rows)} rows x {len(table.columns)} cols")
print(f"Table position: left={table_shape.left}, top={table_shape.top}")
print(f"Table size: width={table_shape.width}, height={table_shape.height}")

# Get column positions
print("\n=== Column Widths ===")
col_positions = [table_shape.left]
for col_idx, col in enumerate(table.columns):
    col_positions.append(col_positions[-1] + col.width)
    print(f"Col {col_idx}: width={col.width}, ends at {col_positions[-1]}")

# Get row positions
print("\n=== Row Heights ===")
row_positions = [table_shape.top]
for row_idx, row in enumerate(table.rows):
    row_positions.append(row_positions[-1] + row.height)
    print(f"Row {row_idx}: height={row.height}, ends at {row_positions[-1]}")

# Map images to cells
print("\n=== Image to Cell Mapping ===")
for img in image_shapes:
    # Use center of image for mapping
    img_center_x = img.left + img.width // 2
    img_center_y = img.top + img.height // 2
    
    # Find column
    col_idx = None
    for i in range(len(col_positions) - 1):
        if col_positions[i] <= img_center_x < col_positions[i + 1]:
            col_idx = i
            break
    
    # Find row
    row_idx = None
    for i in range(len(row_positions) - 1):
        if row_positions[i] <= img_center_y < row_positions[i + 1]:
            row_idx = i
            break
    
    if row_idx is not None and col_idx is not None:
        cell_text = table.cell(row_idx, col_idx).text.strip()[:30] if table.cell(row_idx, col_idx).text else ""
        print(f"'{img.name}' -> Cell[{row_idx},{col_idx}] (text: '{cell_text}')")
    else:
        print(f"'{img.name}' -> Outside table bounds (center: {img_center_x}, {img_center_y})")
