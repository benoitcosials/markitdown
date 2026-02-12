import pptx
from pathlib import Path

pptx_files = list(Path('pptx-test').glob('Kickoff*.pptx'))
pptx_file = pptx_files[0]
prs = pptx.Presentation(str(pptx_file))

print('=' * 80)
print('PROPRIÉTÉS AVANCÉES - DIFFÉRENCIATION ICÔNES vs SCREENSHOTS')
print('=' * 80 + '\n')

for slide_num, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        is_picture = (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE or 
                     (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, 'image')))
        
        if is_picture and hasattr(shape, 'image'):
            img = shape.image
            size_kb = len(img.blob) / 1024
            width_in = shape.width / 914400
            height_in = shape.height / 914400
            
            print(f'\n{"="*60}')
            print(f'Slide {slide_num} - {shape.name}')
            print(f'{"="*60}')
            
            # Basic info
            print(f'\n[BASIC INFO]')
            print(f'  Size: {size_kb:.1f} KB')
            print(f'  Dimensions: {width_in:.2f}" x {height_in:.2f}"')
            print(f'  Aspect Ratio: {width_in/height_in:.2f}')
            
            # Shape name analysis
            print(f'\n[SHAPE NAME ANALYSIS]')
            print(f'  Name: "{shape.name}"')
            if 'Graphique' in shape.name or 'Graph' in shape.name:
                print(f'  → Contains "Graphique/Graph" = Likely ICON')
            elif 'Image' in shape.name or 'Picture' in shape.name:
                print(f'  → Contains "Image/Picture" = Likely PHOTO/SCREENSHOT')
            
            # File properties
            print(f'\n[FILE PROPERTIES]')
            print(f'  Filename: {img.filename}')
            print(f'  Content-Type: {img.content_type}')
            
            # Alt text analysis
            print(f'\n[ALT TEXT]')
            try:
                descr = shape._element._nvXxPr.cNvPr.attrib.get('descr', '')
                if descr:
                    print(f'  Description: "{descr}"')
                    if 'remplissage' in descr.lower() or 'uni' in descr.lower():
                        print(f'  → Generic PowerPoint text = Likely ICON')
                else:
                    print(f'  Description: (empty)')
                    print(f'  → No alt text = Likely IMPORTED IMAGE')
            except:
                print(f'  Description: (unavailable)')
            
            # Image format analysis
            print(f'\n[FORMAT ANALYSIS]')
            # Check if image has transparency
            try:
                from PIL import Image
                import io
                pil_img = Image.open(io.BytesIO(img.blob))
                print(f'  PIL Mode: {pil_img.mode}')
                if pil_img.mode in ['RGBA', 'LA', 'P']:
                    print(f'  → Has transparency = Likely ICON')
                else:
                    print(f'  → No transparency = Likely PHOTO')
                    
                # Check colors
                colors = pil_img.getcolors(maxcolors=1000000)
                if colors:
                    unique_colors = len(colors)
                    print(f'  Unique colors: {unique_colors}')
                    if unique_colors < 100:
                        print(f'  → Few colors (< 100) = Likely ICON/VECTOR')
                    else:
                        print(f'  → Many colors = Likely PHOTO')
            except Exception as e:
                print(f'  PIL Analysis: (unavailable - {e})')
            
            # Final classification
            print(f'\n[CLASSIFICATION HEURISTICS]')
            score_icon = 0
            score_photo = 0
            
            if size_kb < 10:
                print(f'  ✓ Size < 10 KB → +1 ICON')
                score_icon += 1
            elif size_kb > 50:
                print(f'  ✓ Size > 50 KB → +1 PHOTO')
                score_photo += 1
            
            if width_in < 2 and height_in < 2:
                print(f'  ✓ Small dimensions (< 2") → +1 ICON')
                score_icon += 1
            elif width_in > 5 or height_in > 5:
                print(f'  ✓ Large dimensions (> 5") → +1 PHOTO')
                score_photo += 1
            
            if 'Graphique' in shape.name:
                print(f'  ✓ Name contains "Graphique" → +1 ICON')
                score_icon += 1
            elif 'Image' in shape.name:
                print(f'  ✓ Name contains "Image" → +1 PHOTO')
                score_photo += 1
            
            try:
                descr = shape._element._nvXxPr.cNvPr.attrib.get('descr', '')
                if 'remplissage' in descr.lower():
                    print(f'  ✓ Alt text has "remplissage" → +1 ICON')
                    score_icon += 1
                elif not descr:
                    print(f'  ✓ No alt text → +1 PHOTO')
                    score_photo += 1
            except:
                pass
            
            print(f'\n  Score ICON: {score_icon}')
            print(f'  Score PHOTO: {score_photo}')
            
            if score_icon > score_photo:
                print(f'\n  >>> VERDICT: ICON / VECTOR GRAPHIC')
            else:
                print(f'\n  >>> VERDICT: PHOTO / SCREENSHOT')
