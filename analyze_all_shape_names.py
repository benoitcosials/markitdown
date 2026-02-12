import pptx
from pathlib import Path

print('=' * 80)
print('ANALYSE DES NOMS DE SHAPES - TOUS LES PPTX')
print('=' * 80 + '\n')

pptx_files = sorted(Path('pptx-test').glob('*.pptx'))
pptx_files = [f for f in pptx_files if not f.name.startswith('~$')]

all_names = []

for pptx_file in pptx_files:
    print(f'\n{"="*70}')
    print(f'FICHIER: {pptx_file.name}')
    print(f'{"="*70}\n')
    
    try:
        prs = pptx.Presentation(str(pptx_file))
        
        for slide_num, slide in enumerate(prs.slides, 1):
            images_in_slide = []
            
            for shape in slide.shapes:
                is_picture = (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE or 
                             (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, 'image')))
                
                if is_picture and hasattr(shape, 'image'):
                    img = shape.image
                    size_kb = len(img.blob) / 1024
                    width_in = shape.width / 914400
                    height_in = shape.height / 914400
                    
                    # Classification based on size
                    if size_kb < 15:
                        category = 'ICON'
                    elif size_kb > 40:
                        category = 'PHOTO'
                    else:
                        category = 'MIXED'
                    
                    info = {
                        'slide': slide_num,
                        'name': shape.name,
                        'size_kb': size_kb,
                        'width': width_in,
                        'height': height_in,
                        'category': category
                    }
                    
                    images_in_slide.append(info)
                    all_names.append((pptx_file.name, info))
            
            # Print images for this slide
            if images_in_slide:
                for info in images_in_slide:
                    print(f'  Slide {info["slide"]:2d} | {info["category"]:5s} | {info["size_kb"]:6.1f} KB | '
                          f'{info["width"]:5.2f}" x {info["height"]:5.2f}" | "{info["name"]}"')
        
    except Exception as e:
        print(f'  ERREUR: {e}')

# Summary analysis
print(f'\n\n{"="*80}')
print('ANALYSE STATISTIQUE DES NOMS')
print(f'{"="*80}\n')

# Count by name pattern
name_patterns = {}
for pptx_name, info in all_names:
    shape_name = info['name']
    category = info['category']
    
    if shape_name not in name_patterns:
        name_patterns[shape_name] = {'icon': 0, 'photo': 0, 'mixed': 0, 'total': 0}
    
    name_patterns[shape_name][category.lower()] += 1
    name_patterns[shape_name]['total'] += 1

print('Nom du Shape | Total | ICON | PHOTO | MIXED | Verdict')
print('-' * 70)

for name in sorted(name_patterns.keys()):
    stats = name_patterns[name]
    total = stats['total']
    icons = stats['icon']
    photos = stats['photo']
    mixed = stats['mixed']
    
    # Determine verdict
    if icons > photos + mixed:
        verdict = '→ ICON'
    elif photos > icons + mixed:
        verdict = '→ PHOTO'
    else:
        verdict = '→ AMBIGUOUS'
    
    print(f'{name:30s} | {total:5d} | {icons:4d} | {photos:5d} | {mixed:5d} | {verdict}')

print(f'\n{"="*80}')
print('PATTERNS IDENTIFIÉS')
print(f'{"="*80}\n')

# Identify patterns
graphique_count = sum(1 for name in name_patterns if 'Graphique' in name or 'Graph' in name)
image_count = sum(1 for name in name_patterns if 'Image' in name or 'Picture' in name)
other_count = len(name_patterns) - graphique_count - image_count

print(f'Noms contenant "Graphique/Graph": {graphique_count}')
print(f'Noms contenant "Image/Picture": {image_count}')
print(f'Autres noms: {other_count}')

