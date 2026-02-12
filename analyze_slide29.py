from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

pptx_path = r"pptx-test\Standard QA - Design des essais.pptx"
prs = Presentation(pptx_path)

# Slide 29 = index 28 (0-based)
slide_idx = 28
if slide_idx >= len(prs.slides):
    print(f"Erreur: Le fichier ne contient que {len(prs.slides)} slides.")
    exit(1)

slide = prs.slides[slide_idx]
print(f"=== SLIDE 29 - Analyse complete ===")
print(f"Nombre total de shapes: {len(slide.shapes)}")
print()

shape_types = {}
for idx, shape in enumerate(slide.shapes):
    print(f"\n--- Shape {idx} ---")
    print(f"Nom: {shape.name}")
    
    # Type
    shape_type_name = "UNKNOWN"
    try:
        shape_type_name = MSO_SHAPE_TYPE(shape.shape_type).name
    except:
        shape_type_name = str(shape.shape_type)
    
    print(f"Type: {shape_type_name} (valeur: {shape.shape_type})")
    shape_types[shape_type_name] = shape_types.get(shape_type_name, 0) + 1
    
    # Position et dimensions
    print(f"Position: x={shape.left}, y={shape.top}")
    print(f"Dimensions: width={shape.width}, height={shape.height}")
    
    # Texte
    if hasattr(shape, "text"):
        text = shape.text.strip()
        if text:
            preview = text[:80] + "..." if len(text) > 80 else text
            print(f"Texte: {preview}")
    
    # Frame de texte
    if hasattr(shape, "text_frame"):
        print(f"Text frame: {len(shape.text_frame.paragraphs)} paragraphes")
    
    # Image
    if hasattr(shape, "image"):
        try:
            img = shape.image
            size_kb = len(img.blob) / 1024
            print(f"Image: {img.content_type}, {size_kb:.1f} KB")
            print(f"  Extension: {img.ext}")
        except Exception as e:
            pass
    
    # Placeholder (avec gestion exception)
    try:
        if hasattr(shape, "placeholder_format"):
            pf = shape.placeholder_format
            print(f"Placeholder type: {pf.type}")
    except Exception:
        pass
    
    # Shape specifique
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        print("=> PICTURE (image)")
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        print(f"=> GROUP (contient {len(shape.shapes)} shapes)")
    elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        print("=> TEXT_BOX")
    elif shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
        print("=> PLACEHOLDER")
    elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        print(f"=> TABLE ({shape.table.rows} rows x {shape.table.columns} cols)")
    elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
        print("=> CHART")
    elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        print("=> AUTO_SHAPE (forme automatique)")

print(f"\n\n=== RESUME ===")
print(f"Total shapes: {len(slide.shapes)}")
print(f"\nRepartition par type:")
for stype, count in sorted(shape_types.items()):
    print(f"  {stype}: {count}")
