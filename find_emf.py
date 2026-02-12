import pptx
from pathlib import Path

pptx_file = Path('pptx-test').glob('z 015*.pptx')
pptx_file = list(pptx_file)[0]

print(f"Analyse: {pptx_file.name}\n")
prs = pptx.Presentation(str(pptx_file))

for slide_num, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        is_picture = (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE or 
                     (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, 'image')))
        
        if is_picture and hasattr(shape, 'image'):
            img = shape.image
            if '.emf' in (img.filename or '').lower() or 'emf' in (img.content_type or '').lower():
                print(f"Slide {slide_num} - EMF trouvé!")
                print(f"  Nom: {shape.name}")
                print(f"  Filename: {img.filename}")
                print(f"  Content-Type: {img.content_type}")
                print(f"  Size: {len(img.blob) / 1024:.1f} KB")
                print()
