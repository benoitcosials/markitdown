"""
Script détaillé pour analyser les images du slide 19 et vérifier leur classification.
"""
import hashlib
from pathlib import Path

import pptx

# Ouvrir le fichier PPTX
pptx_file = Path("pptx-test/z 015_ARCH-00xx_WS4_Bancaire-b7f2a76a-7655-4a73-a5ab-035aee935269.pptx")
presentation = pptx.Presentation(pptx_file)

# Accéder au slide 19 (index 18, car 0-based)
slide_index = 18
slide = presentation.slides[slide_index]

print("=" * 80)
print(f"ANALYSE DÉTAILLÉE DES IMAGES - SLIDE {slide_index + 1}")
print("=" * 80)

# Fonction de classification (copie depuis _pptx_converter.py)
def classify_image_type(shape) -> str:
    name = shape.name
    
    # Primary classification: shape name patterns
    if 'Graphique' in name or 'Graph' in name:
        return 'icon'
    
    if 'Image' in name or 'Picture' in name:
        return 'photo'
    
    if 'pour une image' in name:
        return 'photo'
    
    # Fallback: file size
    size_kb = len(shape.image.blob) / 1024
    return 'icon' if size_kb < 20 else 'photo'

images = []
for idx, shape in enumerate(slide.shapes):
    if hasattr(shape, 'image'):
        try:
            blob = shape.image.blob
            size_kb = len(blob) / 1024
            classification = classify_image_type(shape)
            md5_hash = hashlib.md5(blob).hexdigest()
            
            images.append({
                'index': idx,
                'name': shape.name,
                'format': shape.image.content_type,
                'size_bytes': len(blob),
                'size_kb': size_kb,
                'width': shape.width,
                'height': shape.height,
                'classification': classification,
                'md5': md5_hash[:12],
                'would_skip': classification == 'icon'
            })
        except Exception as e:
            print(f"Erreur sur shape {idx} ({shape.name}): {e}")

print(f"\nNombre total d'images détectées: {len(images)}\n")
print(f"{'#':<3} {'Nom':<20} {'Format':<15} {'Taille':<12} {'Dimensions':<15} {'Classe':<8} {'MD5':<13} {'Sauté?'}")
print("-" * 110)

for img in images:
    skip_marker = "✗ OUI" if img['would_skip'] else "✓ NON"
    dim_str = f"{img['width']}x{img['height']}"
    size_str = f"{img['size_kb']:.1f} KB"
    
    print(f"{img['index']:<3} {img['name']:<20} {img['format']:<15} {size_str:<12} {dim_str:<15} {img['classification']:<8} {img['md5']:<13} {skip_marker}")

# Vérifier les doublons (déduplication)
print(f"\n{'=' * 80}")
print("VÉRIFICATION DE LA DÉDUPLICATION")
print(f"{'=' * 80}\n")

md5_dict = {}
for img in images:
    if img['md5'] not in md5_dict:
        md5_dict[img['md5']] = []
    md5_dict[img['md5']].append(img['name'])

for md5_hash, names in md5_dict.items():
    if len(names) > 1:
        print(f"✓ MD5 {md5_hash}: {len(names)} doublons → {', '.join(names)}")
    else:
        print(f"  MD5 {md5_hash}: Unique → {names[0]}")

print(f"\n{'=' * 80}")
print("RÉSUMÉ")
print(f"{'=' * 80}\n")

skipped_count = sum(1 for img in images if img['would_skip'])
extracted_count = len(images) - skipped_count
unique_images = len(md5_dict)

print(f"Total images détectées:        {len(images)}")
print(f"Images classées 'icon':        {skipped_count} (sautées par défaut)")
print(f"Images classées 'photo':       {extracted_count} (extraites)")
print(f"Images uniques (après MD5):    {unique_images}")
print("\nPour extraire toutes les images, utiliser: skip_icon_images=False")
