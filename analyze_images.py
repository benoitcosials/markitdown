import pptx
from pathlib import Path

# Load PPTX
pptx_file = Path('pptx-test') / 'Kickoff QA - Essais UAT R1 Bancaire et Trésorerie.pptx'
prs = pptx.Presentation(str(pptx_file))

print('=' * 80)
print('ANALYSE DES IMAGES PPTX - PROPRIÉTÉS DISPONIBLES')
print('=' * 80 + '\n')

for slide_num, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        # Check if it's a picture
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE or \
           (shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, 'image')):
            
            print(f'Slide {slide_num} - {shape.name}')
            print(f'  Type: {shape.shape_type}')
            
            # Image properties
            if hasattr(shape, 'image'):
                img = shape.image
                print(f'  Filename: {img.filename}')
                print(f'  Content-Type: {img.content_type}')
                print(f'  Size (bytes): {len(img.blob)}')
                
                # Shape dimensions
                print(f'  Width: {shape.width} EMUs ({shape.width / 914400:.2f} inches)')
                print(f'  Height: {shape.height} EMUs ({shape.height / 914400:.2f} inches)')
                
                # Check for placeholder type
                if hasattr(shape, 'placeholder_format'):
                    try:
                        ph_type = shape.placeholder_format.type
                        print(f'  Placeholder Type: {ph_type}')
                    except:
                        print(f'  Placeholder Type: Not a placeholder')
                
                # Check shape element for more details
                try:
                    # Picture element
                    pic_elem = shape._element
                    print(f'  Element Tag: {pic_elem.tag}')
                    
                    # Try to get blipFill info
                    if hasattr(pic_elem, 'blipFill'):
                        print(f'  Has blipFill: True')
                    
                    # Check for effects
                    if hasattr(pic_elem, 'spPr'):
                        print(f'  Has spPr (shape properties): True')
                        
                except Exception as e:
                    print(f'  Element inspection error: {e}')
                
                # Check if it has any special attributes
                print(f'  Shape attributes: {[attr for attr in dir(shape) if not attr.startswith("_")]}')
                
            print()
